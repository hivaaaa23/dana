"""استخراج متن از فایل‌های آپلودی (TXT / MD / PDF / DOCX)"""
import io
import os
import shutil
from pathlib import Path

MAX_UPLOAD_BYTES = 10 * 1024 * 1024
SUPPORTED = {".txt", ".md", ".markdown", ".pdf", ".docx"}


# OCR برای PDF اسکن‌شده (اختیاری). نیاز به برنامه‌ی tesseract + داده‌ی زبان فارسی (fas) دارد.
OCR_ENABLED = os.getenv("OCR_ENABLED", "1") not in ("0", "false", "False")
OCR_LANGS = os.getenv("OCR_LANGS", "fas+eng")
OCR_MAX_PAGES = int(os.getenv("OCR_MAX_PAGES", "20"))
OCR_DPI = 200


class ExtractError(ValueError):
    pass


def ocr_available() -> bool:
    if not OCR_ENABLED or shutil.which("tesseract") is None:
        return False
    try:
        import pypdfium2  # noqa: F401
        import pytesseract  # noqa: F401
    except ImportError:
        return False
    return True


def _ocr_pdf(data: bytes) -> str:
    """رندر صفحه‌های PDF به تصویر و OCR (سقف OCR_MAX_PAGES صفحه برای جلوگیری از قفل شدن سرور)"""
    import pypdfium2 as pdfium
    import pytesseract

    try:
        pdf = pdfium.PdfDocument(data)
        n = min(len(pdf), OCR_MAX_PAGES)
        texts = []
        for i in range(n):
            image = pdf[i].render(scale=OCR_DPI / 72).to_pil()
            texts.append(pytesseract.image_to_string(image, lang=OCR_LANGS, timeout=60))
    except pytesseract.TesseractError as e:
        raise ExtractError(
            "OCR اجرا نشد؛ احتمالاً داده‌ی زبان فارسی (fas) برای tesseract نصب نیست."
        ) from e
    except Exception as e:
        raise ExtractError("OCR روی این PDF ناموفق بود.") from e
    return "\n\n".join(texts)


def _decode(data: bytes) -> str:
    # UTF-16 فقط وقتی BOM دارد (بدون BOM تقریباً هر دنباله‌ای را «معتبر» می‌خواند و خروجی آشغال می‌شود)
    if data.startswith((b"\xff\xfe", b"\xfe\xff")):
        return data.decode("utf-16")
    for enc in ("utf-8-sig", "cp1256"):  # cp1256 = ویندوز عربی/فارسی
        try:
            return data.decode(enc)
        except UnicodeError:
            continue
    return data.decode("utf-8", errors="replace")


def _pdf(data: bytes) -> str:
    from pypdf import PdfReader
    from pypdf.errors import PdfReadError

    try:
        reader = PdfReader(io.BytesIO(data))
        if reader.is_encrypted:
            try:
                ok = reader.decrypt("")
            except Exception:
                ok = 0
            if not ok:
                raise ExtractError("این PDF رمزگذاری شده است.")
        return "\n\n".join((page.extract_text() or "") for page in reader.pages)
    except ExtractError:
        raise
    except (PdfReadError, Exception) as e:  # فایل خراب
        raise ExtractError("فایل PDF قابل خواندن نیست.") from e


def _docx(data: bytes) -> str:
    from docx import Document

    try:
        doc = Document(io.BytesIO(data))
    except Exception as e:
        raise ExtractError("فایل DOCX قابل خواندن نیست.") from e
    parts = [p.text for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    return "\n".join(parts)


def extract_text(filename: str, data: bytes) -> str:
    ext = Path(filename or "").suffix.lower()
    if ext not in SUPPORTED:
        raise ExtractError("فرمت پشتیبانی نمی‌شود (TXT، MD، PDF و DOCX مجازند).")
    if len(data) > MAX_UPLOAD_BYTES:
        raise ExtractError("حجم فایل بیش از ۱۰ مگابایت است.")
    text = _pdf(data) if ext == ".pdf" else _docx(data) if ext == ".docx" else _decode(data)
    text = text.replace("\x00", "").strip()
    if not text and ext == ".pdf" and ocr_available():
        text = _ocr_pdf(data).replace("\x00", "").strip()
    if not text:
        hint = (
            " این PDF احتمالاً اسکن‌شده است و OCR روی سرور نصب/فعال نیست."
            if ext == ".pdf" else ""
        )
        raise ExtractError("متنی در فایل پیدا نشد." + hint)
    return text
